"""Excel export of the defect-type x UV test report (your customer table layout)."""
import io
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import ledger
from .db import active_thresholds, now_iso, tx
from .metrics import defect_uv_table

THIN = Side(style="thin", color="7F7F7F")
THICK = Side(style="medium", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
HEAD = PatternFill("solid", fgColor="C6E0B4")
BODY = PatternFill("solid", fgColor="E2EFDA")


def defect_uv_xlsx(conn, by="system"):
    data = defect_uv_table(conn)
    ver, thr = active_thresholds(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "Defect x UV"
    ws["B1"] = "UV Seal Inspection - Test Report by Defect Type and Inspection ROI"
    ws["B1"].font = Font(bold=True, size=13)
    ws["B2"] = f"Generated {now_iso()} - threshold version v{ver} - ground truth from labels (per UV image)"
    ws["B2"].font = Font(italic=True, size=9, color="595959")
    hdr = ["Defect", "INSPECTION ROI", "Categories", "Total No of Images", "Pass", "Fail", "Overkill", "Escape",
           "Not inspectable", "Yield", "Detectability", "Threshold"]
    h = 4
    for k, t in enumerate(hdr):
        c = ws.cell(h, 2 + k, t)
        c.font, c.fill, c.alignment, c.border = Font(bold=True), HEAD, CENTER, BORDER
    ws.merge_cells(start_row=h - 1, start_column=6, end_row=h - 1, end_column=9)
    ws.cell(h - 1, 6, "Prediction").alignment = CENTER
    ws.cell(h - 1, 6).font, ws.cell(h - 1, 6).fill = Font(bold=True), HEAD

    groups = OrderedDict()
    for r in data:
        groups.setdefault(r["defect"], []).append(r)
    row = h + 1
    for defect, items in groups.items():
        start = row
        by_uv = OrderedDict()
        for r in items:
            by_uv.setdefault(r["uv"], []).append(r)
        for uv, lst in by_uv.items():
            uv_start = row
            for r in lst:
                vals = [defect, uv, r["category"], r["total"], r["pass"], r["fail"], r["overkill"], r["escape"],
                        r["not_inspectable"], r["yield"], r["detectability"], thr.get(uv)]
                for k, v in enumerate(vals):
                    c = ws.cell(row, 2 + k, v)
                    c.alignment, c.border, c.fill = CENTER, BORDER, BODY
                    if k in (9, 10):
                        c.number_format = "0.00%"
                row += 1
            if row - uv_start > 1:
                for col in (3, 13):
                    ws.merge_cells(start_row=uv_start, start_column=col, end_row=row - 1, end_column=col)
        ws.merge_cells(start_row=start, start_column=2, end_row=row - 1, end_column=2)
        ws.cell(start, 2).font = Font(bold=True)
        for col in range(2, 2 + len(hdr)):
            c = ws.cell(row - 1, col)
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THICK)
    ws.cell(row + 1, 2, "Overkill = good region predicted Fail; Escape = labelled defect predicted Pass; "
                        "Yield = Pass/Total (Pass rows); Detectability = Fail/Total (Fail rows). "
                        "Not-inspectable images (over-exposed) are not in Total.").font = Font(italic=True, size=9)
    for col, w in zip("ABCDEFGHIJKLM", [2, 24, 15, 12, 14, 9, 9, 10, 9, 14, 10, 13, 11]):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()
    with tx(conn):
        ledger.append(conn, "REPORT_EXPORTED", f"human:{by}", "defect_uv.xlsx",
                      {"sha256": ledger.sha256(content.hex()), "rows": len(data), "threshold_version": ver})
    return content
